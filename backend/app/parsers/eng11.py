from __future__ import annotations

import re
from collections.abc import Sequence
from datetime import date, datetime
from io import BytesIO
from typing import override

import logfire
from docx import Document

from .base import BaseSchedule, DualSchedule, PlannerNote, TextPdfMixin


class Eng11Schedule(DualSchedule, BaseSchedule, TextPdfMixin):
    @classmethod
    @override
    @logfire.instrument("parse eng11 schedule", extract_args=False)
    def from_bytes(
        cls,
        data: bytes,
        course_id: int | None = None,
        *,
        year: int | None = None,
    ) -> Eng11Schedule:
        if data.lstrip().startswith(b"%PDF"):
            return cls.from_pdf_bytes(data, course_id=course_id, year=year)

        document = Document(BytesIO(data))
        assert len(document.tables) == 1, "Expected one syllabus table"
        rows = []
        for row in document.tables[0].rows:
            assert len(row.cells) == 2, "Expected date and assignment columns"
            rows.append(tuple(cell.text.strip() for cell in row.cells))

        assert rows, "Expected a syllabus heading"
        rotation = re.search(r"\b(odd|even)(?:\s+days|s)\b", rows[0][1], re.IGNORECASE)
        assert rotation is not None, "Expected an Odd Days or Even Days heading"
        if len(rows) > 1 and tuple(cell.casefold() for cell in rows[1]) == ("date", "work due"):
            del rows[1]
        notes = cls._parse_section(rows[1:], course_id=course_id, year=year)
        return cls(
            odd_days=notes if rotation[1].lower() == "odd" else [],
            even_days=notes if rotation[1].lower() == "even" else [],
        )

    @classmethod
    def from_pdf_bytes(
        cls,
        data: bytes,
        course_id: int | None = None,
        *,
        year: int | None = None,
    ) -> Eng11Schedule:
        return cls.from_pdf_text(
            cls.extract_text_from_pdf(data), course_id=course_id, year=year
        )

    @classmethod
    @logfire.instrument("parse eng11 PDF schedule", extract_args=False)
    def from_pdf_text(
        cls,
        text: str,
        course_id: int | None = None,
        *,
        year: int | None = None,
    ) -> Eng11Schedule:
        heading = re.search(
            r"\bEng(?:lish)?\.?\s*11\b.{0,200}?\b(Odd|Even)\s+Days\b",
            text,
            re.IGNORECASE | re.DOTALL,
        )
        assert heading is not None, "Expected an English 11 Odd Days or Even Days heading"

        rows = [
            (
                match.group("date"),
                " ".join(
                    re.sub(
                        r"(?im)^.*\bEng(?:lish)?\.?\s*11\b.*"
                        r"\b(?:Odd|Even)\s+Days\b.*$",
                        "",
                        match.group("assignment"),
                    ).split()
                ),
            )
            for match in re.finditer(
                r"^\s*(?P<date>\d{1,2}-[A-Za-z]{3})\s+(?P<assignment>.*?)"
                r"(?=^\s*\d{1,2}-[A-Za-z]{3}\b|\Z)",
                text[heading.end() :],
                re.MULTILINE | re.DOTALL,
            )
        ]
        notes = cls._parse_section(rows, course_id=course_id, year=year)
        return cls(
            odd_days=notes if heading[1].lower() == "odd" else [],
            even_days=notes if heading[1].lower() == "even" else [],
        )

    @staticmethod
    def _parse_section(
        rows: Sequence[tuple[str, ...]],
        course_id: int | None = None,
        *,
        year: int | None = None,
    ) -> list[PlannerNote]:
        if year is None:
            year = date.today().year  # noqa: DTZ011 - calendar year, matching the English 10 parser
        assert rows, "Expected dated assignments"
        start_year = year
        previous_date = None
        notes = []
        for date_text, assignment in rows:
            assert assignment, "Expected an assignment for each date"
            todo_date = datetime.strptime(f"{date_text}-{year}", "%d-%b-%Y").date()
            if previous_date and previous_date.month == 12 and todo_date.month == 1 and year == start_year:
                year += 1
                todo_date = todo_date.replace(year=year)
            if previous_date and todo_date <= previous_date:
                raise ValueError(f"Expected dates in increasing order: {date_text}")
            previous_date = todo_date
            notes.append(
                PlannerNote(
                    todo_date=todo_date,
                    title=" ".join(assignment.split()),
                    course_id=course_id,
                )
            )
        return notes
