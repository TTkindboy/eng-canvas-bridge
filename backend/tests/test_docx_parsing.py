import json
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
import httpx
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from inline_snapshot import snapshot

from app.parsers.base import PlannerNote
from app.dependencies import get_canvas_creds, get_client
from app.parsers.eng11 import Eng11Schedule
from app.routers.pdfs import router

pytestmark = pytest.mark.time_machine(date(2026, 4, 5))


@pytest.fixture
def sample_docx_bytes():
    return (Path(__file__).parent / "data" / "Eng11syllabus2026odds1.docx").read_bytes()


def test_extract_and_parse_full_docx(sample_docx_bytes):
    schedule = Eng11Schedule.from_bytes(sample_docx_bytes, course_id=3565)
    assert schedule == snapshot(
        Eng11Schedule(
            odd_days=[
                PlannerNote(
                    title="Have read the summer reading, Between the World and Me by Ta-Nehisi Coates. Bring the book to class.",
                    todo_date=date(2026, 9, 14),
                    course_id=3565,
                ),
                PlannerNote(
                    title="Find two Coates passages and come up with two questions for discussion in response to Monday's work.",
                    todo_date=date(2026, 9, 16),
                    course_id=3565,
                ),
                PlannerNote(
                    title="Find two Coates passages and come up with two questions for discussion in response to Monday's work. In class: analytic sequence review and practice.",
                    todo_date=date(2026, 9, 18),
                    course_id=3565,
                ),
                PlannerNote(
                    title='Read Emerson\'s "Self-Reliance" through the top of page 4 ("...not quite true..."). Answer the relevant study questions. In class: review of RP1 assignment.',
                    todo_date=date(2026, 9, 24),
                    course_id=3565,
                ),
                PlannerNote(
                    title='Finish "Self-Reliance" and do most of the relevant Emerson study questions.',
                    todo_date=date(2026, 9, 28),
                    course_id=3565,
                ),
                PlannerNote(
                    title="Finish your Emerson study questions. Also: Response Paper 1 due.",
                    todo_date=date(2026, 9, 30),
                    course_id=3565,
                ),
                PlannerNote(
                    title="Read King’s “Letter from Birmingham Jail” (pp. 132-148 in Essays on Civil Disobedience); answer study guide questions 1-15.",
                    todo_date=date(2026, 10, 2),
                    course_id=3565,
                ),
                PlannerNote(
                    title="Complete the King study questions. In class, continue with King’s “Letter.”",
                    todo_date=date(2026, 10, 6),
                    course_id=3565,
                ),
                PlannerNote(
                    title='In Essays on Civil Disobedience, read Thoreau\'s “Civil Disobedience” pp. 22-31 (through "a majority of one already") and answer corresponding study questions.',
                    todo_date=date(2026, 10, 8),
                    course_id=3565,
                ),
                PlannerNote(
                    title="Finish reading Thoreau's “Civil Disobedience” (pp. 31-42) and answer corresponding study questions.",
                    todo_date=date(2026, 10, 14),
                    course_id=3565,
                ),
                PlannerNote(
                    title="No HW! In class: catchup & synthesis discussion of Coates, Emerson, King and Thoreau; read Tracy K. Smith poem. (Bring all of the texts we've read thus far to class.)",
                    todo_date=date(2026, 10, 16),
                    course_id=3565,
                ),
                PlannerNote(
                    title="In-class readings synthesis assignment. Bring all of the texts we've read thus far to class.",
                    todo_date=date(2026, 10, 20),
                    course_id=3565,
                ),
                PlannerNote(
                    title="Study for today's unit test. Bring all of the texts we've read thus far to class.",
                    todo_date=date(2026, 10, 22),
                    course_id=3565,
                ),
            ],
            even_days=[],
        )
    )


def test_explicit_schedule_year(sample_docx_bytes, time_machine):
    time_machine.move_to(date(2027, 4, 5))
    schedule = Eng11Schedule.from_bytes(sample_docx_bytes, year=2026)
    assert schedule.odd_days[0].todo_date == date(2026, 9, 14)
    assert schedule.odd_days[-1].todo_date == date(2026, 10, 22)


def test_even_day_heading(sample_docx_bytes):
    document = Document(BytesIO(sample_docx_bytes))
    document.tables[0].rows[0].cells[
        1
    ].text = "English 11 Syllabus - Even Days - page 1"
    stream = BytesIO()
    document.save(stream)
    schedule = Eng11Schedule.from_bytes(stream.getvalue())
    assert schedule.odd_days == []
    assert len(schedule.even_days) == 13


def test_invalid_date_is_not_silently_skipped(sample_docx_bytes):
    document = Document(BytesIO(sample_docx_bytes))
    document.tables[0].rows[1].cells[0].text = "31-Sep"
    stream = BytesIO()
    document.save(stream)
    with pytest.raises(ValueError):
        Eng11Schedule.from_bytes(stream.getvalue())


@pytest.mark.parametrize(
    "filename, odd_count, even_count",
    [
        ("Eng11syllabus2026odds1.docx", 13, 0),
        ("eng10_sep-oct.pdf", 18, 18),
    ],
)
def test_upload_original_bytes_without_filename(filename, odd_count, even_count):
    app = FastAPI()
    app.include_router(router)
    data = (Path(__file__).parent / "data" / filename).read_bytes()
    with TestClient(app) as client:
        response = client.post(
            "/pdfs/upload", files={"pdf": ("blob", data, "application/octet-stream")}
        )
    assert response.status_code == 200
    schedule = response.json()
    assert len(schedule["odd"]) == odd_count
    assert len(schedule["even"]) == even_count


def test_docx_upload_can_be_submitted_to_canvas(sample_docx_bytes):
    sent_notes = []

    def handle(request):
        assert request.method == "POST"
        assert request.url.path == "/api/v1/planner_notes"
        note = json.loads(request.content)
        sent_notes.append(note)
        return httpx.Response(200, json={**note, "id": len(sent_notes)})

    async def canvas_client():
        async with httpx.AsyncClient(
            base_url="https://canvas.test/api/v1/",
            transport=httpx.MockTransport(handle),
        ) as client:
            yield client

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_client] = canvas_client

    def canvas_auth():
        return {}

    app.dependency_overrides[get_canvas_creds] = canvas_auth
    with TestClient(app) as client:
        upload = client.post("/pdfs/upload", files={"pdf": ("blob", sample_docx_bytes)})
        assert upload.status_code == 200
        schedule = upload.json()
        for day, count in (("odd", 13), ("even", 0)):
            response = client.post(
                "/pdfs/add",
                params={"day": day, "course_id": 3565},
                json=schedule,
            )
            assert response.status_code == 200
            assert len(response.json()) == count
    assert len(sent_notes) == 13
    assert all(note["course_id"] == 3565 for note in sent_notes)
    assert [(note["title"], note["todo_date"]) for note in sent_notes] == [
        (note["title"], note["todo_date"]) for note in schedule["odd"]
    ]
